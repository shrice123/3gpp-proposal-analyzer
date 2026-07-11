from __future__ import annotations

import json
import os
import shutil
import subprocess
import uuid
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor

from .config import settings
from .database import Database, db, utc_now
from .models import call_text_model


DEFAULT_SECTIONS = [
    ("Executive Summary", "Synthesize the central themes and meeting-level implications."),
    ("Proposal Landscape", "Summarize selected TDocs by agenda item and source."),
    ("Technical Analysis", "Explain the key technical changes, dependencies, and alternatives."),
    ("Risks and Open Issues", "Identify unresolved questions, conflicts, and implementation risks."),
    ("Recommendations", "Provide evidence-backed recommendations and next steps."),
]


def create_outline(payload: dict[str, Any], database: Database = db) -> dict[str, Any]:
    report_id = uuid.uuid4().hex
    ids = payload["proposal_ids"]
    placeholders = ",".join("?" for _ in ids)
    proposals = database.query(
        f"""SELECT p.*, a.summary_json,pr.citations_json,pr.visual_evidence_json
            FROM proposals p LEFT JOIN analyses a ON a.proposal_id=p.id
            LEFT JOIN preparations pr ON pr.proposal_id=p.id
            WHERE p.meeting_id=? AND p.id IN ({placeholders}) ORDER BY p.agenda_sort,p.tdoc""",
        tuple([payload["meeting_id"], *ids]),
    )
    summaries = []
    for proposal in proposals:
        summary = json.loads(proposal["summary_json"]) if proposal.get("summary_json") else {}
        summaries.append({
            "tdoc": proposal["tdoc"], "title": proposal["title"], "source": proposal["source"],
            "agenda": proposal["agenda_item"], "summary": summary,
            "citations": json.loads(proposal.get("citations_json") or "[]")[:12],
            "visual_evidence": json.loads(proposal.get("visual_evidence_json") or "[]")[:20],
        })
    outline = _model_outline(payload, summaries, database)
    if outline:
        pass
    elif payload["format"] == "pptx":
        outline = [
            {"title": payload["title"], "role": "opening", "content": "Selected 3GPP proposal analysis"},
            {"title": "The selected set spans the following agenda priorities", "role": "landscape", "content": _landscape(summaries)},
            {"title": "The proposals concentrate on several technical changes", "role": "analysis", "content": _core_changes(summaries)},
            {"title": "The main risks concern dependencies and unresolved assumptions", "role": "risks", "content": _risks(summaries)},
            {"title": "Recommended actions focus review on high-impact dependencies", "role": "recommendation", "content": _recommendations(summaries)},
        ]
    else:
        outline = [
            {"title": title, "purpose": purpose, "content": _section_content(title, summaries)}
            for title, purpose in DEFAULT_SECTIONS
        ]
    now = utc_now()
    database.execute(
        """INSERT INTO reports(id,meeting_id,proposal_ids_json,format,language,template_id,template_mode,title,
           outline_json,status,source_thread_id,source_message_ids_json,context_snapshot_json,created_at,updated_at)
           VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            report_id, payload["meeting_id"], json.dumps(ids), payload["format"], payload.get("language", "en"),
            payload.get("template_id"), payload.get("template_mode", "strict"), payload["title"], json.dumps(outline, ensure_ascii=False),
            "outline", payload.get("source_thread_id"), json.dumps(payload.get("source_message_ids") or []),
            json.dumps(payload.get("conversation_context") or {}, ensure_ascii=False), now, now,
        ),
    )
    return {"id": report_id, "outline": outline, "proposal_count": len(proposals)}


def _model_outline(payload: dict[str, Any], summaries: list[dict[str, Any]], database: Database) -> list[dict[str, Any]]:
    instruction = str(payload.get("instruction") or "").strip()
    if not instruction:
        return []
    language = "Chinese" if payload.get("language") == "zh" else "English"
    unit = "slides" if payload["format"] == "pptx" else "sections"
    prompt = (
        f"Create a concise {language} {payload['format'].upper()} report outline as a JSON array of 4-8 {unit}. "
        "Each item must contain title, content, and role. Focus specifically on the user's instruction. "
        "Preserve TDoc identifiers and do not invent facts. Return JSON only.\n"
        "Use proposal evidence as factual ground truth. Treat user statements as discussion context, not verified facts.\n"
        f"Instruction: {instruction}\nConversation context: {json.dumps(payload.get('conversation_context') or {}, ensure_ascii=False)[:18000]}\n"
        f"Evidence: {json.dumps(summaries, ensure_ascii=False)[:42000]}"
    )
    try:
        raw = call_text_model(
            [{"role": "system", "content": "You prepare evidence-based 3GPP report outlines."},
             {"role": "user", "content": prompt}], database, max_tokens=6000,
        )
        if raw:
            raw = raw.strip().removeprefix("```json").removesuffix("```").strip()
            parsed = json.loads(raw)
            if isinstance(parsed, list) and all(isinstance(item, dict) and item.get("title") for item in parsed):
                return parsed[:12]
    except Exception:
        pass
    return []


def _landscape(items: list[dict[str, Any]]) -> str:
    groups: dict[str, int] = {}
    for item in items:
        groups[item["agenda"]] = groups.get(item["agenda"], 0) + 1
    return "\n".join(f"{agenda}: {count} TDocs" for agenda, count in list(groups.items())[:8])


def _core_changes(items: list[dict[str, Any]]) -> str:
    return "\n".join(
        f"{item['tdoc']}: {item['summary'].get('core_change', item['title'])}" for item in items[:8]
    )


def _risks(items: list[dict[str, Any]]) -> str:
    return "\n".join(f"{item['tdoc']}: {item['summary'].get('risks', 'Requires review')}" for item in items[:8])


def _recommendations(items: list[dict[str, Any]]) -> str:
    return "\n".join(
        f"{item['tdoc']}: {item['summary'].get('recommendation', 'Validate proposal dependencies')}"
        for item in items[:8]
    )


def _section_content(title: str, items: list[dict[str, Any]]) -> str:
    if title == "Proposal Landscape":
        return _landscape(items)
    if title == "Technical Analysis":
        return _core_changes(items)
    if title == "Risks and Open Issues":
        return _risks(items)
    if title == "Recommendations":
        return _recommendations(items)
    return (
        f"This report analyzes {len(items)} selected 3GPP TDocs. "
        "The findings are based on the extracted proposal content and retain source-level traceability."
    )


def inspect_template(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix not in {".docx", ".pptx"}:
        raise ValueError("仅支持 DOCX 或 PPTX 参考文件")
    result: dict[str, Any] = {
        "kind": suffix[1:],
        "file_name": path.name,
        "strict_compatible": True,
        "warnings": [],
    }
    with zipfile.ZipFile(path) as bundle:
        names = bundle.namelist()
        if suffix == ".pptx":
            slides = [name for name in names if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
            layouts = [name for name in names if name.startswith("ppt/slideLayouts/slideLayout") and name.endswith(".xml")]
            placeholders = 0
            for name in slides:
                try:
                    root = ElementTree.fromstring(bundle.read(name))
                    placeholders += sum(1 for node in root.iter() if node.tag.endswith("}ph"))
                except ElementTree.ParseError:
                    pass
            result.update({"slide_count": len(slides), "layout_count": len(layouts), "placeholder_count": placeholders})
            if placeholders == 0:
                result["strict_compatible"] = False
                result["warnings"].append("参考 PPT 没有可识别占位符，需启用风格扩展或准备模板版式")
        else:
            result.update({
                "has_styles": "word/styles.xml" in names,
                "header_count": len([name for name in names if name.startswith("word/header")]),
                "footer_count": len([name for name in names if name.startswith("word/footer")]),
                "has_content_controls": any(b"<w:sdt" in bundle.read(name) for name in names if name.endswith(".xml") and name.startswith("word/")),
                "has_text_placeholders": b"{{" in bundle.read("word/document.xml"),
            })
    return result


def save_template(path: Path, kind: str, name: str, database: Database = db) -> dict[str, Any]:
    template_id = uuid.uuid4().hex
    destination = settings.templates_dir / f"{template_id}{path.suffix.lower()}"
    destination.write_bytes(path.read_bytes())
    inspection = inspect_template(destination)
    database.execute(
        "INSERT INTO templates(id,name,kind,file_path,inspection_json,created_at) VALUES(?,?,?,?,?,?)",
        (template_id, name, kind, str(destination), json.dumps(inspection, ensure_ascii=False), utc_now()),
    )
    return {"id": template_id, "name": name, **inspection}


def render_report(report_id: str, outline: list[dict[str, Any]], database: Database = db) -> dict[str, Any]:
    report = database.one("SELECT * FROM reports WHERE id=?", (report_id,))
    if not report:
        raise ValueError("报告不存在")
    if report["format"] == "docx":
        destination = settings.reports_dir / f"{report_id}.docx"
        _render_docx(report, outline, destination, database)
    else:
        destination = settings.reports_dir / f"{report_id}.pptx"
        _render_pptx(report, outline, destination, database)
    quality = _quality_check(destination)
    preview = _render_office_preview(destination, settings.reports_dir / "previews" / report_id)
    preview_payload = {"quality": quality, "preview_file": str(preview) if preview else None}
    database.execute(
        "UPDATE reports SET outline_json=?,status='completed',file_path=?,preview_json=?,updated_at=? WHERE id=?",
        (json.dumps(outline, ensure_ascii=False), str(destination), json.dumps(preview_payload, ensure_ascii=False), utc_now(), report_id),
    )
    return {
        "id": report_id,
        "file_path": str(destination),
        "file_name": destination.name,
        "quality": quality,
        "preview_available": bool(preview),
    }


def _render_docx(report: dict[str, Any], outline: list[dict[str, Any]], destination: Path, database: Database) -> None:
    template = database.one("SELECT * FROM templates WHERE id=?", (report["template_id"],)) if report.get("template_id") else None
    document = Document(template["file_path"]) if template and template["kind"] == "docx" else Document()
    section = document.sections[0]
    if not template:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
    if template and report.get("template_mode") == "strict":
        replacements = {"{{TITLE}}": report["title"], "{{CONTENT}}": "\n\n".join(str(item.get("content") or item.get("purpose") or "") for item in outline)}
        for index, item in enumerate(outline, start=1):
            replacements[f"{{{{SECTION_{index}_TITLE}}}}"] = str(item.get("title", ""))
            replacements[f"{{{{SECTION_{index}}}}}"] = str(item.get("content") or item.get("purpose") or "")
        replaced = _replace_docx_placeholders(document, replacements)
        if replaced == 0:
            raise ValueError("严格保真模式要求参考 Word 包含 {{TITLE}}、{{CONTENT}} 或分节占位符；也可启用风格扩展")
        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(destination)
        return

    title = document.add_paragraph()
    title.style = document.styles["Title"] if "Title" in [style.name for style in document.styles] else document.styles["Normal"]
    run = title.add_run(report["title"])
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(22, 64, 82)
    metadata = document.add_paragraph("Generated from selected 3GPP meeting proposals")
    metadata.runs[0].font.size = Pt(10)
    metadata.runs[0].font.color.rgb = RGBColor(90, 100, 110)
    for item in outline:
        document.add_heading(item.get("title", "Section"), level=1)
        content = item.get("content") or item.get("purpose", "")
        for line in str(content).splitlines():
            if line.strip():
                document.add_paragraph(line.strip())
    document.add_page_break()
    document.add_heading("Source TDocs", level=1)
    proposal_ids = json.loads(report.get("proposal_ids_json") or "[]")
    if proposal_ids:
        placeholders = ",".join("?" for _ in proposal_ids)
        proposals = database.query(
            f"SELECT tdoc,title,source,agenda_item FROM proposals WHERE meeting_id=? AND id IN ({placeholders}) ORDER BY agenda_sort,tdoc",
            tuple([report["meeting_id"], *proposal_ids]),
        )
    else:
        proposals = database.query("SELECT tdoc,title,source,agenda_item FROM proposals WHERE meeting_id=? ORDER BY agenda_sort,tdoc", (report["meeting_id"],))
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1" if "Light Shading Accent 1" in [style.name for style in document.styles] else "Table Grid"
    for index, label in enumerate(["TDoc", "Title", "Source", "Agenda"]):
        table.rows[0].cells[index].text = label
    for proposal in proposals[:500]:
        cells = table.add_row().cells
        for index, value in enumerate([proposal["tdoc"], proposal["title"], proposal["source"], proposal["agenda_item"]]):
            cells[index].text = str(value)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def _render_pptx(report: dict[str, Any], outline: list[dict[str, Any]], destination: Path, database: Database) -> None:
    template = database.one("SELECT * FROM templates WHERE id=?", (report["template_id"],)) if report.get("template_id") else None
    if template and template["kind"] == "pptx" and report.get("template_mode") == "strict":
        _render_pptx_from_template(Path(template["file_path"]), outline, destination)
        return
    script = Path(os.getenv("PROPOSAL_PPTX_SCRIPT", Path(__file__).resolve().parents[2] / "scripts" / "generate-pptx.mjs"))
    node_path = os.getenv("PROPOSAL_NODE_PATH", "node")
    process_env = os.environ.copy()
    if os.getenv("PROPOSAL_NODE_IS_ELECTRON") == "1":
        process_env["ELECTRON_RUN_AS_NODE"] = "1"
    payload = destination.with_suffix(".json")
    payload.write_text(json.dumps({"title": report["title"], "slides": outline}, ensure_ascii=False), "utf-8")
    result = subprocess.run(
        [node_path, str(script), str(payload), str(destination)],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
        env=process_env,
    )
    payload.unlink(missing_ok=True)
    if result.returncode != 0 or not destination.exists():
        raise RuntimeError(result.stderr or "PPTX 生成失败")


def _replace_docx_placeholders(document: Document, replacements: dict[str, str]) -> int:
    replaced = 0
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        full = paragraph.text
        updated = full
        for marker, value in replacements.items():
            if marker in updated:
                updated = updated.replace(marker, value)
                replaced += 1
        if updated != full:
            for run in paragraph.runs:
                run.text = ""
            if paragraph.runs:
                paragraph.runs[0].text = updated
            else:
                paragraph.add_run(updated)
    return replaced


def _render_pptx_from_template(template: Path, outline: list[dict[str, Any]], destination: Path) -> None:
    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    for prefix, uri in namespaces.items():
        ElementTree.register_namespace(prefix, uri)
    with zipfile.ZipFile(template) as source:
        slide_names = sorted(
            [name for name in source.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")],
            key=lambda name: int("".join(character for character in Path(name).stem if character.isdigit())),
        )
        if len(slide_names) != len(outline):
            raise ValueError(
                f"严格保真模式要求参考 PPT 页数（{len(slide_names)}）与大纲页数（{len(outline)}）一致；也可启用风格扩展"
            )
        modified: dict[str, bytes] = {}
        for slide_name, item in zip(slide_names, outline):
            root = ElementTree.fromstring(source.read(slide_name))
            title_filled = body_filled = False
            for shape in root.findall(".//p:sp", namespaces):
                placeholder = shape.find("./p:nvSpPr/p:nvPr/p:ph", namespaces)
                if placeholder is None:
                    continue
                kind = placeholder.attrib.get("type", "body")
                text_nodes = shape.findall(".//a:t", namespaces)
                if not text_nodes:
                    continue
                if kind in {"title", "ctrTitle"}:
                    text_nodes[0].text = str(item.get("title", ""))
                    for node in text_nodes[1:]:
                        node.text = ""
                    title_filled = True
                elif kind in {"body", "obj", "subTitle"}:
                    text_nodes[0].text = str(item.get("content") or item.get("purpose") or "")
                    for node in text_nodes[1:]:
                        node.text = ""
                    body_filled = True
            if not title_filled:
                raise ValueError(f"参考 PPT 的 {slide_name} 缺少可编辑标题占位符")
            if not body_filled and str(item.get("content") or item.get("purpose") or ""):
                raise ValueError(f"参考 PPT 的 {slide_name} 缺少可编辑正文占位符")
            modified[slide_name] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
        destination.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as output:
            for info in source.infolist():
                output.writestr(info, modified.get(info.filename, source.read(info.filename)))


def _quality_check(path: Path) -> dict[str, Any]:
    warnings: list[str] = []
    with zipfile.ZipFile(path) as bundle:
        if path.suffix.lower() == ".docx":
            document_xml = bundle.read("word/document.xml")
            root = ElementTree.fromstring(document_xml)
            text = "".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
            if len(text.strip()) < 30:
                warnings.append("Word 文档正文内容过少")
        else:
            slide_names = [name for name in bundle.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
            for slide_name in slide_names:
                root = ElementTree.fromstring(bundle.read(slide_name))
                text = "".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
                if not text.strip():
                    warnings.append(f"{slide_name} 没有可见文本")
                for shape in root.iter():
                    if shape.tag.endswith("}ph"):
                        parent_text = text.strip()
                        if not parent_text:
                            warnings.append(f"{slide_name} 存在未填充占位符")
                            break
    return {"passed": not warnings, "warnings": warnings}


def _render_office_preview(path: Path, output_dir: Path) -> Path | None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        return None
    output_dir.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [executable, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(path)],
        capture_output=True,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = output_dir / f"{path.stem}.pdf"
    return pdf if result.returncode == 0 and pdf.exists() else None
