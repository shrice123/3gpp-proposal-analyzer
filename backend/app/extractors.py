from __future__ import annotations

import shutil
import subprocess
import zipfile
import json
import hashlib
import multiprocessing
import re
import time
import uuid
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {".doc", ".docx", ".ppt", ".pptx", ".pdf", ".xls", ".xlsx", ".txt"}
MAX_ARCHIVE_FILES = 200
MAX_ARCHIVE_BYTES = 750 * 1024 * 1024


def safe_extract(archive: Path, destination: Path) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []
    total = 0
    with zipfile.ZipFile(archive) as bundle:
        members = bundle.infolist()
        if len(members) > MAX_ARCHIVE_FILES:
            raise ValueError("压缩包文件数量超过安全限制")
        for member in members:
            total += member.file_size
            if total > MAX_ARCHIVE_BYTES:
                raise ValueError("压缩包展开大小超过安全限制")
            target = (destination / member.filename).resolve()
            if destination.resolve() not in target.parents and target != destination.resolve():
                raise ValueError("压缩包包含不安全路径")
            if member.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with bundle.open(member) as source, target.open("wb") as output:
                shutil.copyfileobj(source, output)
            files.append(target)
    return files


def _xml_text(archive: Path, prefixes: tuple[str, ...]) -> tuple[str, int]:
    parts: list[str] = []
    visual_count = 0
    with zipfile.ZipFile(archive) as bundle:
        for name in sorted(bundle.namelist()):
            if name.startswith(prefixes) and name.endswith(".xml"):
                try:
                    root = ElementTree.fromstring(bundle.read(name))
                    text = " ".join(node.text.strip() for node in root.iter() if node.text and node.text.strip())
                    if text:
                        parts.append(text)
                except ElementTree.ParseError:
                    continue
            if "/media/" in name and not name.endswith("/"):
                visual_count += 1
    return "\n\n".join(parts), visual_count


def extract_office_media(archive: Path, output_dir: Path, media_prefix: str) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    extracted: list[Path] = []
    allowed = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
    with zipfile.ZipFile(archive) as bundle:
        for index, name in enumerate(sorted(bundle.namelist())):
            if not name.startswith(media_prefix) or Path(name).suffix.lower() not in allowed:
                continue
            data = bundle.read(name)
            if len(data) < 2048 or len(data) > 20 * 1024 * 1024:
                continue
            destination = output_dir / f"{index:03d}-{Path(name).name}"
            destination.write_bytes(data)
            extracted.append(destination)
    return extracted


def extract_office_visual_evidence(path: Path) -> list[dict]:
    evidence: list[dict] = []
    with zipfile.ZipFile(path) as bundle:
        for name in sorted(bundle.namelist()):
            if name.startswith(("ppt/charts/chart", "word/charts/chart", "xl/charts/chart")) and name.endswith(".xml"):
                try:
                    root = ElementTree.fromstring(bundle.read(name))
                except ElementTree.ParseError:
                    continue
                labels = [node.text.strip() for node in root.iter() if node.text and node.text.strip()]
                if labels:
                    chart_text = lambda suffix: [
                        text.text.strip() for node in root.iter() if node.tag.endswith(suffix)
                        for text in node.iter() if text.text and text.text.strip()
                    ]
                    evidence.append({
                        "kind": "office_chart", "artifact": path.name, "part": name,
                        "chart_type": next((node.tag.rsplit("}", 1)[-1].removesuffix("Chart") for node in root.iter() if node.tag.endswith("Chart")), "chart"),
                        "series": chart_text("}ser")[:80], "categories": chart_text("}cat")[:160],
                        "values": chart_text("}val")[:320],
                        "summary": " | ".join(labels[:240]), "confidence": 0.95,
                        "resolved": True, "resolution_mode": "local_structure",
                    })
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                try:
                    root = ElementTree.fromstring(bundle.read(name))
                except ElementTree.ParseError:
                    continue
                shapes: list[dict] = []
                connectors: list[dict] = []
                for node in root.iter():
                    if node.tag.endswith("}cxnSp"):
                        endpoints = [
                            {"kind": child.tag.rsplit("}", 1)[-1], "shape_id": child.attrib.get("id"), "connection_index": child.attrib.get("idx")}
                            for child in node.iter() if child.tag.endswith(("}stCxn", "}endCxn"))
                        ]
                        connectors.append({"endpoints": endpoints})
                    if not node.tag.endswith("}sp"):
                        continue
                    text = " ".join(child.text.strip() for child in node.iter() if child.tag.endswith("}t") and child.text and child.text.strip())
                    if text:
                        offset = next((child for child in node.iter() if child.tag.endswith("}off")), None)
                        extent = next((child for child in node.iter() if child.tag.endswith("}ext")), None)
                        non_visual = next((child for child in node.iter() if child.tag.endswith("}cNvPr")), None)
                        shapes.append({
                            "id": non_visual.attrib.get("id") if non_visual is not None else None,
                            "text": text[:500],
                            "position": {"x": offset.attrib.get("x"), "y": offset.attrib.get("y")} if offset is not None else None,
                            "size": {"width": extent.attrib.get("cx"), "height": extent.attrib.get("cy")} if extent is not None else None,
                        })
                if connectors and len(shapes) >= 2:
                    slide_number = int(re.search(r"slide(\d+)", name).group(1)) if re.search(r"slide(\d+)", name) else None
                    evidence.append({
                        "kind": "diagram_structure", "artifact": path.name, "part": name,
                        "page_or_slide": slide_number, "shapes": shapes[:80], "connectors": connectors[:160],
                        "summary": json.dumps({"shapes": shapes[:80], "connectors": connectors[:160]}, ensure_ascii=False),
                        "confidence": 0.8, "resolved": True, "resolution_mode": "local_structure",
                    })
    return evidence


def _ocr_image(path: Path) -> tuple[str, list[dict]]:
    executable = shutil.which("tesseract")
    if not executable:
        return "", []
    try:
        result = subprocess.run(
            [executable, str(path), "stdout", "-l", "eng", "tsv"], capture_output=True,
            text=True, timeout=90, check=False,
        )
        if result.returncode != 0:
            return "", []
        words: list[str] = []
        blocks: list[dict] = []
        for line in result.stdout.splitlines()[1:]:
            columns = line.split("\t", 11)
            if len(columns) != 12 or not columns[11].strip():
                continue
            try:
                confidence = float(columns[10]) / 100
                box = {"x": int(columns[6]), "y": int(columns[7]), "width": int(columns[8]), "height": int(columns[9])}
            except ValueError:
                continue
            words.append(columns[11].strip())
            blocks.append({"text": columns[11].strip(), "confidence": round(confidence, 3), **box})
        return " ".join(words)[:8000], blocks[:1000]
    except Exception:
        return "", []


def image_visual_evidence(path: Path, artifact: str) -> dict | None:
    if not path.exists() or path.stat().st_size < 2048:
        return None
    ocr, blocks = _ocr_image(path)
    if len(ocr) < 24 and path.stat().st_size < 50 * 1024:
        return None
    page_match = re.search(r"page-(\d+)", path.name)
    return {
        "kind": "image_ocr", "artifact": artifact, "asset_path": str(path),
        "asset_sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "summary": ocr,
        "page_or_slide": int(page_match.group(1)) if page_match else None,
        "ocr_blocks": blocks, "confidence": 0.45 if len(ocr) >= 24 else 0.15,
        "resolved": False, "resolution_mode": "local_ocr" if blocks else "local_unresolved",
        "unresolved_reason": "位图无法通过本地处理可靠恢复图形关系" if blocks else "位图缺少可可靠提取的结构或文字",
    }


def extract_docx(path: Path) -> tuple[str, int]:
    return _xml_text(path, ("word/document.xml", "word/header", "word/footer", "word/footnotes"))


def extract_pptx(path: Path) -> tuple[str, int]:
    return _xml_text(path, ("ppt/slides/slide", "ppt/notesSlides/notesSlide", "ppt/charts/chart"))


def extract_xlsx(path: Path) -> tuple[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    output: list[str] = []
    for sheet in workbook.worksheets:
        output.append(f"Sheet: {sheet.title}")
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_index >= 5000:
                break
            values = [str(value) for value in row if value not in (None, "")]
            if values:
                output.append(" | ".join(values))
    workbook.close()
    return "\n".join(output), 0


def extract_pdf(path: Path) -> tuple[str, int]:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages), sum(len(page.images) for page in reader.pages)


def extract_pdf_media(path: Path, output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(str(path))
    extracted: list[Path] = []
    for page_index, page in enumerate(reader.pages):
        for image_index, image in enumerate(page.images):
            suffix = Path(image.name).suffix.lower() or ".png"
            if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
                continue
            data = image.data
            if len(data) < 2048 or len(data) > 20 * 1024 * 1024:
                continue
            destination = output_dir / f"page-{page_index + 1}-image-{image_index + 1}{suffix}"
            destination.write_bytes(data)
            extracted.append(destination)
    return extracted


def convert_legacy(path: Path, output_dir: Path) -> Path | None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        return None
    output_dir.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [executable, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(path)],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    candidate = output_dir / f"{path.stem}.pdf"
    return candidate if result.returncode == 0 and candidate.exists() else None


def extract_document(path: Path, conversion_dir: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        return {"file": path.name, "text": "", "visual_count": 0, "status": "unsupported"}
    try:
        visual_files: list[Path] = []
        visual_evidence: list[dict] = []
        if suffix == ".docx":
            text, visuals = extract_docx(path)
            visual_files = extract_office_media(path, conversion_dir / "media" / path.stem, "word/media/")
            visual_evidence = extract_office_visual_evidence(path)
        elif suffix == ".pptx":
            text, visuals = extract_pptx(path)
            visual_files = extract_office_media(path, conversion_dir / "media" / path.stem, "ppt/media/")
            visual_evidence = extract_office_visual_evidence(path)
        elif suffix == ".xlsx":
            text, visuals = extract_xlsx(path)
            visual_evidence = extract_office_visual_evidence(path)
        elif suffix == ".pdf":
            text, visuals = extract_pdf(path)
            visual_files = extract_pdf_media(path, conversion_dir / "media" / path.stem)
        elif suffix == ".txt":
            text, visuals = path.read_text("utf-8", errors="replace"), 0
        else:
            converted = convert_legacy(path, conversion_dir)
            if not converted:
                return {
                    "file": path.name,
                    "text": "",
                    "visual_count": 0,
                    "status": "runtime_required",
                }
            text, visuals = extract_pdf(converted)
            visual_files = extract_pdf_media(converted, conversion_dir / "media" / path.stem)
        visual_evidence.extend(
            item for item in (image_visual_evidence(visual, path.name) for visual in visual_files) if item
        )
        return {
            "file": path.name,
            "text": text,
            "visual_count": max(visuals, len(visual_files)),
            "visual_files": [str(item) for item in visual_files],
            "visual_evidence": visual_evidence,
            "status": "ok",
        }
    except Exception as exc:
        return {"file": path.name, "text": "", "visual_count": 0, "status": "failed", "error": str(exc)}


def _isolated_extract_worker(path: str, conversion_dir: str, result_path: str) -> None:
    result = extract_document(Path(path), Path(conversion_dir))
    Path(result_path).write_text(json.dumps(result, ensure_ascii=False), "utf-8")


def extract_document_isolated(path: Path, conversion_dir: Path, timeout: int = 300, cancelled=None) -> dict:
    """Run parser code outside the job runner so a wedged Office/PDF parser can be terminated."""
    conversion_dir.mkdir(parents=True, exist_ok=True)
    result_path = conversion_dir / f".extract-result-{uuid.uuid4().hex}.json"
    context = multiprocessing.get_context("spawn")
    process = context.Process(target=_isolated_extract_worker, args=(str(path), str(conversion_dir), str(result_path)))
    process.start()
    deadline = time.monotonic() + timeout
    try:
        while process.is_alive() and time.monotonic() < deadline:
            if cancelled and cancelled():
                process.terminate()
                process.join(5)
                raise RuntimeError("文稿提取已取消")
            process.join(.25)
        if process.is_alive():
            process.terminate()
            process.join(5)
            return {"file": path.name, "text": "", "visual_count": 0, "status": "failed", "error": "文稿提取超过 5 分钟，已终止"}
        if process.exitcode != 0 or not result_path.exists():
            return {"file": path.name, "text": "", "visual_count": 0, "status": "failed", "error": "文稿提取进程异常退出"}
        return json.loads(result_path.read_text("utf-8"))
    finally:
        result_path.unlink(missing_ok=True)
