from __future__ import annotations

import hashlib
import re
import shutil
import uuid
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import quote, unquote, urljoin, urlparse

import httpx
from openpyxl import load_workbook

from .config import settings
from .database import Database, db, utc_now
from .source_normalization import apply_meeting_parent_names, canonical_source_name, extract_primary_source


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36"
)
INDEX_PATTERN = re.compile(r"TDoc_List_Meeting_.*\.xlsx$", re.IGNORECASE)
TDOC_PATTERN = re.compile(r"^[A-Za-z0-9]+-[A-Za-z0-9-]+$")
SUPPORTED_DOCUMENT_SUFFIXES = {".zip", ".docx", ".pptx", ".pdf", ".xlsx"}
DEFAULT_FTP_URL = "https://www.3gpp.org/ftp/tsg_sa/wg2_arch/"


def primary_source(source: str) -> str:
    return extract_primary_source(source)


class LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() != "a":
            return
        for key, value in attrs:
            if key.lower() == "href" and value:
                # Keep percent-encoding intact. 3GPP index filenames contain '#';
                # decoding it before urljoin would incorrectly turn the suffix
                # into a URL fragment and make the XLSX undiscoverable.
                self.links.append(value)


class DirectoryTableParser(HTMLParser):
    """Read the public 3GPP directory table without depending on its styling."""

    def __init__(self) -> None:
        super().__init__()
        self.rows: list[dict[str, Any]] = []
        self._row: dict[str, Any] | None = None
        self._cell: list[str] | None = None
        self._cells: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_dict = dict(attrs)
        if tag.lower() == "tr":
            self._row, self._cells = {"href": "", "name": ""}, []
        elif tag.lower() == "td" and self._row is not None:
            self._cell = []
        elif tag.lower() == "a" and self._row is not None:
            href = attrs_dict.get("href")
            if href and not href.startswith("?") and not href.startswith("#"):
                self._row["href"] = href

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "td" and self._cell is not None:
            self._cells.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag.lower() == "tr" and self._row is not None:
            if self._row.get("href"):
                values = [value for value in self._cells if value and value.casefold() != "icon"]
                name = unquote(urlparse(self._row["href"]).path.rstrip("/").split("/")[-1])
                self._row.update({
                    "name": name,
                    "modified": next((v for v in values if re.match(r"^\d{4}/\d{2}/\d{2}", v)), ""),
                    "size_text": next((v for v in values if re.search(r"\b(?:bytes?|kb|mb|gb)\b", v, re.I)), ""),
                })
                self.rows.append(self._row)
            self._row, self._cell, self._cells = None, None, []


def _parse_size(value: str) -> int | None:
    match = re.search(r"([\d.,]+)\s*(bytes?|kb|mb|gb)", value or "", re.I)
    if not match:
        return None
    number = match.group(1).replace(".", "").replace(",", ".")
    try:
        amount = float(number)
    except ValueError:
        return None
    multiplier = {"byte": 1, "bytes": 1, "kb": 1024, "mb": 1024 ** 2, "gb": 1024 ** 3}[match.group(2).casefold()]
    return int(amount * multiplier)


def _breadcrumbs(url: str) -> list[dict[str, str]]:
    parsed = urlparse(url)
    parts = [unquote(part) for part in parsed.path.split("/") if part]
    crumbs = [{"label": "www.3gpp.org", "url": "https://www.3gpp.org/"}]
    path = ""
    for part in parts:
        path += "/" + quote(part, safe="")
        crumbs.append({"label": part, "url": f"https://www.3gpp.org{path}/"})
    return crumbs


def validate_3gpp_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme.lower() != "https":
        raise ValueError("仅支持 HTTPS 的 3GPP FTP 目录链接")
    if parsed.hostname not in {"www.3gpp.org", "3gpp.org"}:
        raise ValueError("链接域名必须是 www.3gpp.org")
    if not parsed.path.lower().startswith("/ftp/"):
        raise ValueError("链接必须位于 3GPP 的 /ftp/ 目录下")
    if parsed.username or parsed.password or parsed.port:
        raise ValueError("链接中不能包含凭据或自定义端口")
    return url.rstrip("/") + "/"


def meeting_name_from_url(url: str) -> str:
    parts = [part for part in urlparse(url).path.split("/") if part]
    if parts and parts[-1].lower() == "docs":
        parts.pop()
    return parts[-1] if parts else "3GPP Meeting"


def http_client(referer: str | None = None) -> httpx.Client:
    headers = {"User-Agent": USER_AGENT, "Accept": "*/*"}
    if referer:
        headers["Referer"] = referer
    return httpx.Client(headers=headers, follow_redirects=True, timeout=60)


def list_directory(url: str) -> list[str]:
    validated = validate_3gpp_url(url)
    with http_client(validated) as client:
        response = client.get(validated)
        response.raise_for_status()
    content_type = str(getattr(response, "headers", {}).get("content-type", ""))
    if content_type and "html" not in content_type.casefold():
        raise ValueError("链接必须指向 3GPP FTP 文件夹，不能直接指向文件")
    parser = LinkParser()
    parser.feed(response.text)
    return [urljoin(validated, link) for link in parser.links]


def list_directory_entries(url: str) -> list[dict[str, Any]]:
    validated = validate_3gpp_url(url)
    with http_client(validated) as client:
        response = client.get(validated)
        response.raise_for_status()
    content_type = str(getattr(response, "headers", {}).get("content-type", ""))
    if content_type and "html" not in content_type.casefold():
        raise ValueError("链接必须指向 3GPP FTP 文件夹，不能直接指向文件")
    parser = DirectoryTableParser()
    parser.feed(response.text)
    entries: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in parser.rows:
        absolute = urljoin(validated, row["href"])
        parsed = urlparse(absolute)
        if parsed.hostname not in {"www.3gpp.org", "3gpp.org"} or not parsed.path.lower().startswith("/ftp/"):
            continue
        clean_url = absolute.split("?", 1)[0].split("#", 1)[0]
        if clean_url in seen or clean_url.rstrip("/") == validated.rstrip("/"):
            continue
        seen.add(clean_url)
        suffix = Path(row["name"]).suffix.casefold()
        is_folder = not row.get("size_text")
        if is_folder:
            clean_url = clean_url.rstrip("/") + "/"
        entries.append({
            "kind": "folder" if is_folder else "file",
            "name": row["name"], "url": clean_url,
            "modified": row.get("modified") or "", "size": _parse_size(row.get("size_text") or ""),
            "extension": suffix.lstrip("."), "analyzable": suffix in SUPPORTED_DOCUMENT_SUFFIXES,
        })
    return entries


def find_index_url(links: list[str]) -> str:
    candidates = [link for link in links if INDEX_PATTERN.search(unquote(urlparse(link).path).split("/")[-1])]
    if not candidates:
        raise ValueError("目录中未找到 TDoc_List_Meeting_*.xlsx，请手动上传索引表")
    return sorted(candidates)[-1]


def download_file(url: str, destination: Path, referer: str | None = None, cancelled: Any | None = None) -> dict[str, Any]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    partial = destination.with_suffix(destination.suffix + ".part")
    headers: dict[str, str] = {}
    if partial.exists() and partial.stat().st_size:
        headers["Range"] = f"bytes={partial.stat().st_size}-"
    with http_client(referer) as client:
        with client.stream("GET", url, headers=headers) as response:
            if response.status_code == 416:
                partial.replace(destination)
                return {"bytes": destination.stat().st_size, "etag": "", "last_modified": ""}
            response.raise_for_status()
            mode = "ab" if response.status_code == 206 else "wb"
            with partial.open(mode) as handle:
                for chunk in response.iter_bytes(1024 * 256):
                    if cancelled and cancelled():
                        raise RuntimeError("下载已取消")
                    handle.write(chunk)
            metadata = {
                "bytes": partial.stat().st_size,
                "etag": response.headers.get("etag", ""),
                "last_modified": response.headers.get("last-modified", ""),
            }
    partial.replace(destination)
    return metadata


HEADER_ALIASES = {
    "tdoc": {"tdoc", "t doc", "tdoc number", "document number"},
    "title": {"title", "tdoc title", "document title"},
    "source": {"source", "company", "organisation", "organization"},
    "agenda_item": {"agenda item", "agenda", "agenda_item"},
    "agenda_description": {"agenda item description", "agenda description", "agenda_item_description"},
    "agenda_sort": {"agenda item sort order", "agenda sort order"},
    "status": {"tdoc status", "status"},
    "abstract": {"abstract", "summary"},
}


def _normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _field_indexes(headers: list[Any]) -> dict[str, int]:
    normalized = [_normalize(value) for value in headers]
    result: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for index, header in enumerate(normalized):
            if header in aliases:
                result[field] = index
                break
    missing = [field for field in ("tdoc", "title", "source", "agenda_item") if field not in result]
    if missing:
        raise ValueError(f"Excel 缺少必要字段：{', '.join(missing)}")
    return result


def parse_index(path: Path, listing_links: list[str], source_url: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["TDoc_List"] if "TDoc_List" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    indexes = _field_indexes(headers)
    link_map = {unquote(urlparse(link).path).split("/")[-1].casefold(): link for link in listing_links}

    proposals: list[dict[str, Any]] = []
    for row in rows:
        tdoc = str(row[indexes["tdoc"]] or "").strip()
        if not tdoc or not TDOC_PATTERN.match(tdoc):
            continue
        zip_name = f"{tdoc}.zip"
        zip_url = link_map.get(zip_name.casefold(), urljoin(source_url, quote(zip_name)))
        agenda_sort = row[indexes["agenda_sort"]] if "agenda_sort" in indexes else 999999
        try:
            agenda_sort_value = float(agenda_sort)
        except (TypeError, ValueError):
            agenda_sort_value = 999999
        source = str(row[indexes["source"]] or "Unknown").strip() or "Unknown"
        proposals.append(
            {
                "tdoc": tdoc,
                "title": str(row[indexes["title"]] or "").strip(),
                "source": source,
                "primary_source": primary_source(source),
                "canonical_source": canonical_source_name(source),
                "agenda_item": str(row[indexes["agenda_item"]] or "Unassigned").strip(),
                "agenda_description": str(
                    row[indexes["agenda_description"]] if "agenda_description" in indexes else ""
                ).strip(),
                "agenda_sort": agenda_sort_value,
                "status": str(row[indexes["status"]] if "status" in indexes else "").strip(),
                "abstract": str(row[indexes["abstract"]] if "abstract" in indexes else "").strip(),
                "zip_url": zip_url,
            }
        )
    workbook.close()
    return proposals


def browse_folder(source_url: str, database: Database = db) -> dict[str, Any]:
    """Browse one FTP level and materialize analyzable files as proposal records."""
    source_url = validate_3gpp_url(source_url or DEFAULT_FTP_URL)
    entries = list_directory_entries(source_url)
    files = [item for item in entries if item["kind"] == "file"]
    folders = [item for item in entries if item["kind"] == "folder"]
    index_entry = next((item for item in files if INDEX_PATTERN.search(item["name"])), None)
    index_status = "absent"
    index_error = ""
    indexed: list[dict[str, Any]] = []
    if index_entry:
        index_status = "ready"
        seed = hashlib.sha256(index_entry["url"].encode("utf-8")).hexdigest()[:20]
        index_path = settings.cache_dir / "meetings" / seed / "index.xlsx"
        try:
            download_file(index_entry["url"], index_path, source_url)
            indexed = parse_index(index_path, [item["url"] for item in files], source_url)
        except Exception as exc:
            index_status, index_error, indexed = "failed", str(exc), []

    meeting_id = hashlib.sha256(source_url.encode("utf-8")).hexdigest()[:24]
    now = utc_now()
    by_name = {item["name"].casefold(): item for item in files}
    records: list[dict[str, Any]] = []
    consumed: set[str] = set()
    for item in indexed:
        remote = by_name.get(f"{item['tdoc']}.zip".casefold())
        if remote:
            consumed.add(remote["url"])
        file_url = remote["url"] if remote else item["zip_url"]
        records.append({
            **item, "file_url": file_url, "file_name": remote["name"] if remote else f"{item['tdoc']}.zip",
            "file_kind": "zip", "file_size": remote.get("size") if remote else None,
            "file_modified": remote.get("modified") if remote else None, "metadata_source": "index",
            "analyzable": bool(remote),
        })
    for remote in files:
        if remote["url"] in consumed or not remote["analyzable"]:
            continue
        suffix = Path(remote["name"]).suffix.casefold()
        stem = Path(remote["name"]).stem
        records.append({
            "tdoc": stem, "title": remote["name"], "source": "Unknown", "primary_source": "Unknown",
            "canonical_source": "Unknown", "agenda_item": "Unassigned", "agenda_description": "",
            "agenda_sort": 999999, "status": "available", "abstract": "", "zip_url": remote["url"],
            "file_url": remote["url"], "file_name": remote["name"], "file_kind": suffix.lstrip("."),
            "file_size": remote["size"], "file_modified": remote["modified"], "metadata_source": "directory",
            "analyzable": True,
        })
    with database.transaction() as connection:
        connection.execute(
            """INSERT INTO meetings(id,name,source_url,index_file,proposal_count,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name,index_file=excluded.index_file,
               proposal_count=excluded.proposal_count,updated_at=excluded.updated_at""",
            (meeting_id, meeting_name_from_url(source_url), source_url,
             str(index_path) if index_entry and index_status == "ready" else None, len(records), now, now),
        )
        current_ids: set[str] = set()
        for item in records:
            proposal_id = uuid.uuid5(uuid.NAMESPACE_URL, f"3gpp-file:{item['file_url']}").hex
            connection.execute(
                """INSERT INTO proposals(id,meeting_id,tdoc,title,source,primary_source,canonical_source,
                   agenda_item,agenda_description,agenda_sort,status,abstract,zip_url,file_url,file_name,file_kind,
                   file_size,file_modified,metadata_source,analyzable)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(meeting_id,tdoc) DO UPDATE SET title=excluded.title,source=excluded.source,
                   primary_source=excluded.primary_source,canonical_source=excluded.canonical_source,
                   agenda_item=excluded.agenda_item,agenda_description=excluded.agenda_description,
                   agenda_sort=excluded.agenda_sort,status=excluded.status,abstract=excluded.abstract,
                   zip_url=excluded.zip_url,file_url=excluded.file_url,file_name=excluded.file_name,
                   file_kind=excluded.file_kind,file_size=excluded.file_size,file_modified=excluded.file_modified,
                   metadata_source=excluded.metadata_source,analyzable=excluded.analyzable""",
                (proposal_id, meeting_id, item["tdoc"], item["title"], item["source"], item["primary_source"],
                 item["canonical_source"], item["agenda_item"], item["agenda_description"], item["agenda_sort"],
                 item["status"], item["abstract"], item["zip_url"], item["file_url"], item["file_name"],
                 item["file_kind"], item["file_size"], item["file_modified"], item["metadata_source"],
                 int(item["analyzable"])),
            )
            stored = connection.execute(
                "SELECT id FROM proposals WHERE meeting_id=? AND tdoc=?", (meeting_id, item["tdoc"])
            ).fetchone()
            current_ids.add(stored["id"] if stored else proposal_id)
        old_rows = connection.execute("SELECT id FROM proposals WHERE meeting_id=?", (meeting_id,)).fetchall()
        for row in old_rows:
            if row["id"] not in current_ids:
                connection.execute("DELETE FROM proposals WHERE id=?", (row["id"],))
    proposal_rows = database.query("SELECT * FROM proposals WHERE meeting_id=? ORDER BY agenda_sort,tdoc", (meeting_id,))
    raw_files = []
    proposal_by_url = {row["file_url"]: row for row in proposal_rows}
    for item in files:
        raw_files.append({**item, "proposal": proposal_by_url.get(item["url"])})
    return {
        "workspace": {"id": meeting_id, "name": meeting_name_from_url(source_url), "source_url": source_url,
                      "proposal_count": len(proposal_rows)},
        "breadcrumbs": _breadcrumbs(source_url), "folders": folders, "files": raw_files,
        "proposals": proposal_rows, "index_status": index_status, "index_error": index_error,
        "facets": {
            "agendas": ["all", *sorted({row["agenda_item"] for row in proposal_rows if row["agenda_item"] != "Unassigned"})],
            "sources": ["all", *sorted({row["canonical_source"] for row in proposal_rows if row["canonical_source"] != "Unknown"})],
        },
    }


def import_meeting(source_url: str, database: Database = db, index_path: Path | None = None) -> dict[str, Any]:
    source_url = validate_3gpp_url(source_url)
    links = list_directory(source_url)
    if index_path is None:
        index_url = find_index_url(links)
        seed = hashlib.sha256(source_url.encode("utf-8")).hexdigest()[:12]
        index_path = settings.cache_dir / "meetings" / seed / "index.xlsx"
        download_file(index_url, index_path, source_url)
    proposals = parse_index(index_path, links, source_url)
    user_aliases = {
        row["alias_key"]: row["canonical_name"]
        for row in database.query("SELECT alias_key,canonical_name FROM source_aliases")
    }
    for item in proposals:
        item["canonical_source"] = canonical_source_name(item["primary_source"], user_aliases)
    apply_meeting_parent_names(proposals, user_aliases)
    meeting_id = hashlib.sha256(source_url.encode("utf-8")).hexdigest()[:24]
    now = utc_now()
    with database.transaction() as connection:
        connection.execute(
            """INSERT INTO meetings(id,name,source_url,index_file,proposal_count,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?)
               ON CONFLICT(id) DO UPDATE SET name=excluded.name,index_file=excluded.index_file,
               proposal_count=excluded.proposal_count,updated_at=excluded.updated_at""",
            (
                meeting_id,
                meeting_name_from_url(source_url),
                source_url,
                str(index_path),
                len(proposals),
                now,
                now,
            ),
        )
        existing = {
            row["tdoc"]: dict(row)
            for row in connection.execute("SELECT * FROM proposals WHERE meeting_id=?", (meeting_id,)).fetchall()
        }
        changed_items = [
            (uuid.uuid5(uuid.NAMESPACE_URL, f"{meeting_id}:{item['tdoc']}").hex, item["tdoc"])
            for item in proposals
            if item["tdoc"] in existing
            and (existing[item["tdoc"]]["zip_url"] != item["zip_url"] or existing[item["tdoc"]]["title"] != item["title"])
        ]
        connection.executemany(
            """INSERT INTO proposals(
               id,meeting_id,tdoc,title,source,primary_source,canonical_source,agenda_item,agenda_description,agenda_sort,status,abstract,zip_url
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(meeting_id,tdoc) DO UPDATE SET title=excluded.title,source=excluded.source,
               primary_source=excluded.primary_source,canonical_source=excluded.canonical_source,agenda_item=excluded.agenda_item,
               agenda_description=excluded.agenda_description,agenda_sort=excluded.agenda_sort,
               status=excluded.status,abstract=excluded.abstract,zip_url=excluded.zip_url""",
            [
                (
                    uuid.uuid5(uuid.NAMESPACE_URL, f"{meeting_id}:{item['tdoc']}").hex,
                    meeting_id,
                    item["tdoc"],
                    item["title"],
                    item["source"],
                    item["primary_source"],
                    item["canonical_source"],
                    item["agenda_item"],
                    item["agenda_description"],
                    item["agenda_sort"],
                    item["status"],
                    item["abstract"],
                    item["zip_url"],
                )
                for item in proposals
            ],
        )
        current_tdocs = {item["tdoc"] for item in proposals}
        for old_tdoc, old in existing.items():
            if old_tdoc not in current_tdocs:
                connection.execute("DELETE FROM proposals WHERE id=?", (old["id"],))
        for proposal_id, _ in changed_items:
            connection.execute("DELETE FROM preparations WHERE proposal_id=?", (proposal_id,))
            connection.execute("DELETE FROM analyses WHERE proposal_id=?", (proposal_id,))
            connection.execute("DELETE FROM analysis_fts WHERE proposal_id=?", (proposal_id,))
            connection.execute(
                "UPDATE proposals SET download_state='idle',preparation_state='idle',preparation_error=NULL,analysis_state='idle' WHERE id=?",
                (proposal_id,),
            )
    for _, tdoc in changed_items:
        (settings.cache_dir / "zips" / f"{tdoc}.zip").unlink(missing_ok=True)
        shutil.rmtree(settings.cache_dir / "extracted" / tdoc, ignore_errors=True)
        shutil.rmtree(settings.cache_dir / "converted" / tdoc, ignore_errors=True)
    return {
        "id": meeting_id,
        "name": meeting_name_from_url(source_url),
        "source_url": source_url,
        "proposal_count": len(proposals),
    }


def list_proposals(
    meeting_id: str,
    *,
    agenda: list[str] | str | None = None,
    source: list[str] | str | None = None,
    search: str = "",
    limit: int = 500,
    offset: int = 0,
    database: Database = db,
) -> dict[str, Any]:
    agendas = [agenda] if isinstance(agenda, str) else list(agenda or [])
    sources_selected = [source] if isinstance(source, str) else list(source or [])
    agendas = [item for item in agendas if item and item != "all"]
    sources_selected = [item for item in sources_selected if item and item != "all"]
    def filters(*, include_agenda: bool, include_source: bool) -> tuple[str, list[Any]]:
        clauses = ["meeting_id=?"]
        values: list[Any] = [meeting_id]
        if include_agenda and agendas:
            clauses.append(f"agenda_item IN ({','.join('?' for _ in agendas)})")
            values.extend(agendas)
        if include_source and sources_selected:
            clauses.append(f"canonical_source COLLATE NOCASE IN ({','.join('?' for _ in sources_selected)})")
            values.extend(sources_selected)
        if search:
            clauses.append("(tdoc LIKE ? OR title LIKE ? OR abstract LIKE ?)")
            term = f"%{search}%"
            values.extend([term, term, term])
        return " AND ".join(clauses), values

    where, params = filters(include_agenda=True, include_source=True)
    total_row = database.one(f"SELECT COUNT(*) AS count FROM proposals WHERE {where}", tuple(params))
    proposals = database.query(
        f"""SELECT * FROM proposals WHERE {where}
            ORDER BY agenda_sort, tdoc LIMIT ? OFFSET ?""",
        tuple(params + [min(limit, 2500), max(offset, 0)]),
    )
    agenda_where, agenda_params = filters(include_agenda=False, include_source=True)
    agenda_rows = database.query(
        f"""SELECT agenda_item AS value,MIN(agenda_sort) AS sort FROM proposals
            WHERE {agenda_where} GROUP BY agenda_item ORDER BY sort,value""", tuple(agenda_params),
    )
    source_where, source_params = filters(include_agenda=True, include_source=False)
    sources = database.query(
        f"""SELECT canonical_source AS value FROM proposals WHERE {source_where}
            GROUP BY canonical_source COLLATE NOCASE ORDER BY value COLLATE NOCASE""", tuple(source_params),
    )
    return {
        "items": proposals,
        "total": total_row["count"] if total_row else 0,
        "facets": {
            "agendas": ["all"] + [item["value"] for item in agenda_rows],
            "sources": ["all"] + [item["value"] for item in sources],
        },
    }
